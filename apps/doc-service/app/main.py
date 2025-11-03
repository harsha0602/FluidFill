import asyncio
import html
import json
import logging
import os
import re
from collections import OrderedDict
from io import BytesIO, UnsupportedOperation
from typing import Any, Dict, Iterable, List, Optional, Tuple
from zipfile import BadZipFile

from dotenv import load_dotenv
from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from fastapi import FastAPI, File, HTTPException, UploadFile
from pydantic import BaseModel

from app.ai_studio import call_gemini_json
from app.models import SchemaResponse

load_dotenv()

app = FastAPI(title="FluidFill Doc Service")
logger = logging.getLogger(__name__)

_DEPRECATED_ENV_VARS = [
    "GOOGLE_APPLICATION_CREDENTIALS",
    "GOOGLE_CLOUD_PROJECT",
    "GOOGLE_API_KEY",
    "GOOGLE_GENAI_MODEL",
    "GOOGLE_GENAI_USE_VERTEXAI",
]
for _deprecated in _DEPRECATED_ENV_VARS:
    if _deprecated in os.environ:
        os.environ.pop(_deprecated, None)

AI_STUDIO_MODEL = os.getenv("AI_STUDIO_MODEL", "gemini-1.5-flash")

logger.info("AI Studio model: %s", AI_STUDIO_MODEL)


def _get_ai_studio_api_key() -> Optional[str]:
    return os.getenv("AI_STUDIO_API_KEY")


if not _get_ai_studio_api_key():
    logger.warning("AI_STUDIO_API_KEY not configured; /schema endpoint will return 500 until set.")


class ParsedPlaceholder(BaseModel):
    key: str
    label: str
    occurrences: int


class ParseResult(BaseModel):
    document_id: str
    placeholders: List[ParsedPlaceholder]


class Placeholder(BaseModel):
    key: str
    label: str
    occurrences: int
    example_context: Optional[str] = None


PLACEHOLDER_PATTERN = re.compile(r"\[([^\]]+)\]|(_{3,})")

HTML_STYLES = """
<style>
body { background: #1a1b1f; color: #f4f4f5; font-family: 'Inter', sans-serif; margin: 0; padding: 1.5rem; }
.doc-body { max-width: 60rem; margin: 0 auto; }
.placeholder { color: #9d76dd; background: rgba(157,118,221,0.1); border-radius: 4px; padding: 0 3px; }
p { margin-bottom: 0.75rem; line-height: 1.6; }
table { width: 100%; border-collapse: collapse; margin-bottom: 1.25rem; }
td, th { border: 1px solid rgba(255,255,255,0.12); padding: 0.5rem; vertical-align: top; }
.empty-line { min-height: 1rem; }
</style>
""".strip()


def _to_snake_case(value: str) -> str:
    cleaned = re.sub(r"[^\w\s]", " ", value.strip())
    parts = [part.lower() for part in cleaned.split() if part]
    return "_".join(parts)


def _iter_block_items(parent: Any) -> Iterable[Any]:
    """Yield paragraphs and tables in document order for the given parent."""
    if isinstance(parent, DocumentType):
        parent_element = parent.element.body  # type: ignore[attr-defined]
    else:
        parent_element = getattr(parent, "_element", None)
    if parent_element is None:
        return []

    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _highlight_placeholders(text: str) -> Tuple[str, int]:
    """Return HTML with highlighted placeholders and the count of replacements."""
    if not text:
        return "", 0

    result_parts: List[str] = []
    last_index = 0
    placeholder_count = 0

    for match in PLACEHOLDER_PATTERN.finditer(text):
        result_parts.append(html.escape(text[last_index:match.start()]))
        raw_label = match.group(1)
        if raw_label is not None:
            label = " ".join(raw_label.split())
            key = _to_snake_case(label)
            if not key:
                result_parts.append(html.escape(match.group(0)))
            else:
                placeholder_count += 1
                placeholder_literal = html.escape(match.group(0))
                result_parts.append(
                    f'<span class="placeholder" data-key="{html.escape(key)}">{placeholder_literal}</span>'
                )
        else:
            underscores = match.group(2)
            if underscores:
                placeholder_count += 1
                placeholder_literal = html.escape(underscores)
                result_parts.append(
                    '<span class="placeholder" data-key="blank">'
                    f"{placeholder_literal}</span>"
                )
        last_index = match.end()

    result_parts.append(html.escape(text[last_index:]))
    rendered = "".join(result_parts).replace("\n", "<br />")
    return rendered, placeholder_count


def _render_paragraph(paragraph: Paragraph) -> Tuple[str, int]:
    text = paragraph.text or ""
    rendered, placeholder_count = _highlight_placeholders(text)
    if not rendered.strip():
        rendered = "&nbsp;"
        css_class = ' class="empty-line"'
    else:
        css_class = ""
    return f"<p{css_class}>{rendered}</p>", placeholder_count


def _render_table(table: Table) -> Tuple[str, int, int]:
    placeholder_total = 0
    paragraph_total = 0
    rows_html: List[str] = []

    for row in table.rows:
        cells_html: List[str] = []
        for cell in row.cells:
            cell_parts: List[str] = []
            cell_placeholder_count, cell_paragraphs = _render_blocks(cell, cell_parts)
            placeholder_total += cell_placeholder_count
            paragraph_total += cell_paragraphs
            cell_content = "".join(cell_parts).strip() or "&nbsp;"
            cells_html.append(f"<td>{cell_content}</td>")
        rows_html.append(f"<tr>{''.join(cells_html)}</tr>")

    table_html = f'<table class="doc-table">{"".join(rows_html)}</table>'
    return table_html, placeholder_total, paragraph_total


def _render_blocks(parent: Any, container: List[str]) -> Tuple[int, int]:
    placeholder_total = 0
    paragraph_total = 0

    for block in _iter_block_items(parent):
        if isinstance(block, Paragraph):
            paragraph_html, ph_count = _render_paragraph(block)
            container.append(paragraph_html)
            paragraph_total += 1
            placeholder_total += ph_count
        elif isinstance(block, Table):
            table_html, table_placeholders, table_paragraphs = _render_table(block)
            container.append(table_html)
            paragraph_total += table_paragraphs
            placeholder_total += table_placeholders

    return placeholder_total, paragraph_total


def _iter_all_paragraphs(parent: Any) -> Iterable[Paragraph]:
    for block in _iter_block_items(parent):
        if isinstance(block, Paragraph):
            yield block
        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    yield from _iter_all_paragraphs(cell)


def extract_placeholders(document: DocumentType) -> List[ParsedPlaceholder]:
    placeholders: OrderedDict[str, Dict[str, Any]] = OrderedDict()

    for paragraph in _iter_all_paragraphs(document):
        if not paragraph.runs:
            continue
        text = "".join(run.text for run in paragraph.runs if run.text)
        if not text:
            continue

        for match in PLACEHOLDER_PATTERN.finditer(text):
            raw_label = match.group(1)
            if raw_label is not None:
                label = " ".join(raw_label.split())
                if not label:
                    continue
                key = _to_snake_case(label)
                if not key:
                    continue
            else:
                label = "Blank"
                key = "blank"

            placeholder_entry = placeholders.setdefault(
                key, {"label": label, "occurrences": 0}
            )
            placeholder_entry["occurrences"] += 1

    return [
        ParsedPlaceholder(key=key, label=data["label"], occurrences=data["occurrences"])
        for key, data in placeholders.items()
    ]


def render_docx_to_html(file: UploadFile) -> str:
    file_stream = getattr(file, "file", None)
    if file_stream is None:
        raise ValueError("Missing file stream")

    try:
        file_stream.seek(0)
    except (AttributeError, UnsupportedOperation):
        pass

    contents = file_stream.read()
    if isinstance(contents, str):
        contents = contents.encode("utf-8")
    if not isinstance(contents, (bytes, bytearray)):
        raise ValueError("Unable to read uploaded document stream")

    try:
        document = Document(BytesIO(contents))
    except (BadZipFile, ValueError, KeyError) as exc:
        raise ValueError("Invalid DOCX file") from exc

    html_parts: List[str] = []
    placeholder_total, paragraph_total = _render_blocks(document, html_parts)

    document_name = file.filename or "document"
    logger.info(
        "Rendered HTML preview for %s with %d paragraphs and %d placeholders",
        document_name,
        paragraph_total,
        placeholder_total,
    )

    html_content = "".join(html_parts) or "<p>&nbsp;</p>"
    full_html = (
        "<!DOCTYPE html><html><head>"
        f"{HTML_STYLES}"
        "</head><body><div class=\"doc-body\">"
        f"{html_content}"
        "</div></body></html>"
    )

    try:
        file_stream.seek(0)
    except (AttributeError, UnsupportedOperation):
        pass

    return full_html


def build_schema_prompt(placeholders: List[Placeholder]) -> str:
    placeholder_payload = json.dumps(
        [placeholder.model_dump(exclude_none=True) for placeholder in placeholders],
        indent=2,
    )
    return (
        "You are generating a structured form schema to help users fill a legal SAFE agreement.\n"
        "Return JSON only that conforms to this schema:\n"
        'SchemaResponse = {"groups": [{"id": str, "title": str, "description"?: str, "fields": ['
        '{"key": str, "label": str, "type": "text"|"email"|"phone"|"date"|"number"|"multiline"|"select", '
        '"required": bool, "help"?: str, "repeat_group"?: str, "targets": [str]}]}]}\n'
        "Guidelines:\n"
        "- Group fields logically by topic (e.g., Company, Investor, Economics).\n"
        "- Set field.targets to every placeholder key this field should populate; use the key values exactly as provided.\n"
        "- When one answer fills multiple placeholders, set repeat_group to a stable snake_case id (e.g., \"company_name\").\n"
        "- Use occurrences and context to decide when placeholders share the same field.\n"
        "- Choose accurate field.type values: email for emails, date for dates, phone for phone numbers, number for numeric values, multiline for long free-form responses, select when a discrete choice is implied.\n"
        "- Provide concise help text only when it adds clarity.\n"
        "- Keep the JSON compact (aim for <400 tokens) and return nothing else.\n"
        "- Do not include markdown fences, comments, or explanations outside the JSON.\n"
        "\n"
        "Placeholders:\n"
        f"{placeholder_payload}\n"
    )


@app.get("/health")
async def health():
    return {"ok": True, "service": "doc-service"}


@app.post("/parse", response_model=ParseResult)
async def parse(file: UploadFile = File(...)):
    contents = await file.read()
    document_id = file.filename or "document"
    try:
        document = Document(BytesIO(contents))
    except (BadZipFile, ValueError, KeyError) as exc:
        logger.exception("Unable to open DOCX for %s", document_id)
        raise HTTPException(status_code=400, detail="Invalid DOCX file") from exc

    placeholder_models = extract_placeholders(document)

    logger.info(
        "Parsed %d placeholder types (%d total occurrences) from %s",
        len(placeholder_models),
        sum(placeholder.occurrences for placeholder in placeholder_models),
        document_id,
    )

    return ParseResult(
        document_id=document_id,
        placeholders=placeholder_models,
    )


@app.post("/to_html")
async def to_html(file: UploadFile = File(...)):
    try:
        html_content = render_docx_to_html(file)
    except ValueError as exc:
        logger.exception("Unable to render HTML for %s", file.filename or "document")
        raise HTTPException(status_code=400, detail="Invalid DOCX file") from exc
    except Exception as exc:  # pragma: no cover - unexpected errors
        logger.exception("Unexpected error rendering HTML for %s", file.filename or "document")
        raise HTTPException(status_code=500, detail="Failed to render HTML") from exc

    return {"html": html_content}


@app.post("/schema")
async def schema_endpoint(payload: Dict[str, Any]):
    if not _get_ai_studio_api_key():
        logger.error("AI_STUDIO_API_KEY not configured; cannot generate schema")
        raise HTTPException(status_code=500, detail="AI_STUDIO_API_KEY not configured")

    items = payload.get("placeholders", []) if isinstance(payload, dict) else []
    placeholders: List[Placeholder] = []
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            continue
        raw_label = item.get("label") or item.get("name") or ""
        label = str(raw_label).strip()
        raw_key = item.get("key")
        key = str(raw_key).strip() if isinstance(raw_key, str) else ""
        if not key and label:
            key = _to_snake_case(label)
        if not label and key:
            label = key.replace("_", " ").title()
        if not key:
            key = f"field_{index}"
        if not label:
            label = key.replace("_", " ").title()
        occurrences_value = item.get("occurrences", 1)
        try:
            occurrences = max(int(occurrences_value), 1)
        except (TypeError, ValueError):
            occurrences = 1
        placeholders.append(
            Placeholder(
                key=key,
                label=label,
                occurrences=occurrences,
                example_context=item.get("example_context"),
            )
        )

    prompt = build_schema_prompt(placeholders)
    logger.info("Requesting schema from Google AI Studio for %d placeholders", len(placeholders))

    try:
        raw = await asyncio.to_thread(call_gemini_json, prompt, AI_STUDIO_MODEL)
    except Exception:
        logger.exception("AI Studio schema generation failed")
        raise HTTPException(status_code=502, detail="schema_generation_failed")

    try:
        schema_payload = json.loads(raw)
        schema = SchemaResponse.model_validate(schema_payload)
    except Exception as exc:  # pragma: no cover - defensive logging
        logger.error("Invalid schema JSON: %s ...raw=%s", exc, raw[:600])
        raise HTTPException(status_code=502, detail="schema_generation_failed") from exc

    logger.info("Received schema with %d groups from Google AI Studio", len(schema.groups))
    return schema.model_dump()


@app.get("/dev/ai-studio-ping")
def ai_studio_ping():
    if os.getenv("ENABLE_DEV_ROUTES", "false").lower() != "true":
        raise HTTPException(status_code=404, detail="Not found")
    prompt = 'Return JSON: {"ok": true}'
    raw = call_gemini_json(prompt, AI_STUDIO_MODEL)
    return {"raw": raw[:2000]}
